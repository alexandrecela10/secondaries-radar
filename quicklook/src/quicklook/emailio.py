"""Turn raw email bytes into plain content: headers, body text, attachment tables."""
from __future__ import annotations

import csv
import datetime as dt
import email
import hashlib
import io
from dataclasses import dataclass, field
from email import policy
from email.utils import parsedate_to_datetime


@dataclass
class Table:
    filename: str
    rows: list[list[str]]  # header row first

    def as_csv(self) -> str:
        buf = io.StringIO()
        csv.writer(buf).writerows(self.rows)
        return buf.getvalue()


@dataclass
class EmailContent:
    message_id: str
    subject: str
    sender: str
    date: dt.datetime  # when the broker sent it; our point-in-time "as of"
    body: str
    tables: list[Table] = field(default_factory=list)

    def as_text(self) -> str:
        """What a reader (or the model) sees: subject, body, then each table as CSV."""
        parts = [f"Subject: {self.subject}", "", self.body]
        for t in self.tables:
            parts += ["", f"[Attachment {t.filename}]", t.as_csv()]
        return "\n".join(parts)


def normalise_message_id(raw: str | None, body: bytes) -> str:
    if raw:
        return raw.strip().strip("<>").strip().lower()
    # No Message-ID: fall back to a content hash so the same bytes still dedupe.
    return "sha256-" + hashlib.sha256(body).hexdigest()[:32]


def _xlsx_rows(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
    # Brokers put a title above the header. The header is the first wide row.
    for i, r in enumerate(rows):
        if sum(1 for c in r if c.strip()) >= 5:
            return [x for x in rows[i:] if any(c.strip() for c in x)]
    return rows


def read_email(raw: bytes) -> EmailContent:
    msg = email.message_from_bytes(raw, policy=policy.default)
    date_hdr = msg.get("Date")
    try:
        date = parsedate_to_datetime(date_hdr) if date_hdr else None
    except (TypeError, ValueError):
        date = None
    if date is None:
        date = dt.datetime(1970, 1, 1, tzinfo=dt.timezone.utc)
    if date.tzinfo is None:
        date = date.replace(tzinfo=dt.timezone.utc)
    date = date.astimezone(dt.timezone.utc)
    body_part = msg.get_body(preferencelist=("plain", "html"))
    body = body_part.get_content() if body_part is not None else ""
    tables = []
    for part in msg.iter_attachments():
        name = part.get_filename() or "attachment"
        data = part.get_payload(decode=True) or b""
        if name.lower().endswith(".csv"):
            text = data.decode("utf-8-sig", errors="replace")
            tables.append(Table(name, [r for r in csv.reader(io.StringIO(text)) if r]))
        elif name.lower().endswith(".xlsx"):
            tables.append(Table(name, _xlsx_rows(data)))
    return EmailContent(
        message_id=normalise_message_id(msg.get("Message-ID"), raw),
        subject=str(msg.get("Subject", "")).replace("\n", " ").strip(),
        sender=str(msg.get("From", "")),
        date=date,
        body=body,
        tables=tables,
    )
